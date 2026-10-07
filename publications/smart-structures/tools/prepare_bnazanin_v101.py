#!/usr/bin/env python3
"""Prepare Smart Structures v1.0.1 B Nazanin typography revision.

This script changes Persian typography to the font family name "B Nazanin"
without bundling or redistributing proprietary font files. A licensed local
installation is required for faithful rendering and final PDF generation.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

PERSIAN_RE = re.compile(r"[\u0600-\u06FF]")
OLD_PERSIAN_FONTS = {"Vazirmatn", "Sahel", "Shabnam"}
FONT = "B Nazanin"


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{slot}"), name)


def set_style_font(style, size: float, bold: bool | None = None) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{slot}"), FONT)


def prepare(source: Path, destination: Path) -> None:
    doc = Document(source)

    specs = {
        "Normal": (12.0, None),
        "List Bullet": (12.0, None),
        "List Number": (12.0, None),
        "Heading 1": (17.5, True),
        "Heading 2": (14.0, True),
        "Caption": (10.5, True),
    }
    for name, (size, bold) in specs.items():
        set_style_font(doc.styles[name], size, bold)

    doc.styles["Normal"].paragraph_format.line_spacing = 1.12
    doc.styles["Normal"].paragraph_format.space_after = Pt(4.5)
    for name in ("List Bullet", "List Number"):
        doc.styles[name].paragraph_format.line_spacing = 1.10
        doc.styles[name].paragraph_format.space_after = Pt(2.5)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(4.5)
    doc.styles["Caption"].paragraph_format.space_before = Pt(3)
    doc.styles["Caption"].paragraph_format.space_after = Pt(4)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if not run.text or run.font.name == "Cambria Math":
                continue
            if PERSIAN_RE.search(run.text) or run.font.name in OLD_PERSIAN_FONTS:
                explicit = run.font.size.pt if run.font.size else None
                set_run_font(run, FONT, explicit if explicit and explicit >= 16 else None)
                if not (explicit and explicit >= 16):
                    run.font.size = None

    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.03
                    paragraph.paragraph_format.space_after = Pt(1)
                    for run in paragraph.runs:
                        if run.font.name == "Cambria Math" or not run.text.strip():
                            continue
                        set_run_font(run, FONT, 11.25 if (run.bold or row_index == 0) else 10.75)

    # The v1.0.0 source intentionally forced the closing quotation onto a new page.
    # For v1.0.1 keep the closing block with the bibliography so the book stays 40 pages.
    closing = next(
        (p for p in doc.paragraphs if "سازه‌ای که می‌آموزد، باید پیش از هر چیز ایمن بماند" in p.text),
        None,
    )
    if closing is not None:
        ppr = closing._p.get_or_add_pPr()
        page_break = ppr.find(qn("w:pageBreakBefore"))
        if page_break is not None:
            ppr.remove(page_break)
        closing.paragraph_format.space_before = Pt(28)
        closing.paragraph_format.space_after = Pt(8)

    props = doc.core_properties
    props.subject = "Smart Structures and Seismic Response Control — v1.0.1 B Nazanin typography revision"
    props.comments = (
        "Persian typography mapped to B Nazanin. The font is referenced but not embedded. "
        "A licensed local installation is required for faithful rendering and final PDF publication."
    )
    props.revision = max(int(props.revision or 0), 4)

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    prepare(args.input, args.output)
    print(f"Prepared {args.output}")


if __name__ == "__main__":
    main()
